"""Normalize a local game workbook without modifying or publishing its source.

Read cached Excel values only. Repair ingredient prices by merchant NAME, never
row position. This is not an Excel calculation engine or a profitability model.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path

FIELDS = {
    '食材': ['Type', 'Name', 'Level', 'Sell Price', 'Buy Price', 'Tags', 'Owned'],
    '料理': ['Name', 'Owned', 'Cookware', 'Sell Price', 'Tags', 'Required Ingredients', 'Forbidden Ingredients'],
    '稀客': ['Name', 'Budget', 'Preferred Tags', 'Disliked Tags', 'Beverage Tags', 'Spawn Areas', 'Type'],
    '食材merchant_data': ['Type', 'Name', 'Total Daily Quantity', 'commission',
        *[field for i in range(1, 6) for field in (f'Buy Price{i}', f'Daily Quantity{i}')],
        'Average Buy Price'],
}


def finite_number(value):
    if value is None or isinstance(value, bool) or value == '':
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def names(value):
    return [part.strip() for part in re.split(r'[、,，;/；|\n\r\t]+', str(value or '')) if part.strip()]


def named_rows(rows, sheet):
    result = {}
    for row in rows:
        name = str(row.get('Name') or '').strip()
        if not name:
            continue
        if name in result:
            raise ValueError(f'{sheet}: duplicate Name {name!r}')
        row['Name'] = name
        result[name] = row
    return result


def normalize_sheets(raw):
    """Return normalized data and explicit corrections; useful independently of XLSX."""
    missing = set(FIELDS) - raw.keys()
    if missing:
        raise ValueError('Missing sheets: ' + ', '.join(sorted(missing)))
    report = {'price_corrections': [], 'unknown_prices': [], 'disabled_recipes': [], 'warnings': []}
    sheets = {}
    for sheet, fields in FIELDS.items():
        cleaned = []
        for source_row in raw[sheet]:
            row = {key: source_row.get(key) for key in fields if key in source_row}
            for key, value in row.items():
                if isinstance(value, str) and value.startswith(('#DIV/0!', '#VALUE!', '#N/A', '#REF!', '#NUM!')):
                    row[key] = None
                elif isinstance(value, float) and not math.isfinite(value):
                    row[key] = None
            cleaned.append(row)
        sheets[sheet] = list(named_rows(cleaned, sheet).values())

    merchants = named_rows(sheets['食材merchant_data'], '食材merchant_data')
    ingredients = named_rows(sheets['食材'], '食材')
    for name, ingredient in ingredients.items():
        if name not in merchants:
            raise ValueError(f'No merchant supply row for {name!r}')
        merchant = merchants[name]
        quantity = finite_number(merchant.get('Total Daily Quantity'))
        if quantity is None or quantity < 0:
            raise ValueError(f'Invalid or uncached Total Daily Quantity for {name!r}')
        merchant['Total Daily Quantity'] = quantity
        old = finite_number(ingredient.get('Buy Price'))
        price = finite_number(merchant.get('Average Buy Price'))
        if price is not None and price < 0:
            raise ValueError(f'Negative Average Buy Price for {name!r}')
        merchant['Average Buy Price'] = price
        ingredient['Buy Price'] = price
        if price is None:
            report['unknown_prices'].append(name)
        if old != price:
            report['price_corrections'].append({'Name': name, 'old': old, 'new': price,
                'method': 'Match 食材merchant_data.Average Buy Price by Name'})

    for recipe in sheets['料理']:
        base = names(recipe.get('Required Ingredients'))
        unknown = sorted(set(base) - ingredients.keys())
        reasons = []
        if not base:
            reasons.append('No concrete base recipe')
        if unknown:
            reasons.append('Unknown or symbolic ingredients: ' + ', '.join(unknown))
        if len(base) > 5:
            reasons.append('More than five ingredient units')
        if not str(recipe.get('Cookware') or '').strip():
            reasons.append('No cookware')
        if recipe['Name'] in {'黑暗物质', '糟糕料理'}:
            reasons.append('Non-plannable special recipe')
        if reasons:
            recipe['Owned'] = 0
            report['disabled_recipes'].append({'Name': recipe['Name'], 'reasons': reasons})
        price = finite_number(recipe.get('Sell Price'))
        if price is None or price < 0:
            recipe['Owned'] = 0
            recipe['Sell Price'] = 0
            report['disabled_recipes'].append({'Name': recipe['Name'], 'reasons': ['Missing or invalid sell price']})
    report['row_counts'] = {key: len(rows) for key, rows in sheets.items()}
    if report['unknown_prices']:
        report['warnings'].append('Unknown acquisition prices remain null; do not interpret them as free ingredients.')
    report['warnings'].append('Cached Excel values retained. Save/recalculate your workbook in a spreadsheet app if formulas have no cache.')
    report['warnings'].append('Supply and price are author estimates, not a current-game accuracy guarantee.')
    return {'format': 'mystia-planner-data/v1', 'metadata': {
        'title': '本地导入的游戏数据', 'kind': 'local-game-data',
        'distribution': 'Local use; review upstream rights before redistributing.',
        'normalization': 'Price lookup repaired by merchant Name; source workbook untouched.',
    }, 'sheets': sheets}, report


def read_workbook(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit('XLSX import requires: python -m pip install -r requirements.txt') from exc
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        raw = {}
        for name in FIELDS:
            if name not in workbook.sheetnames:
                raise ValueError(f'Missing sheet: {name}')
            rows = iter(workbook[name].values)
            header = [str(value).strip() if value is not None else '' for value in next(rows)]
            raw[name] = [dict(zip(header, row)) for row in rows]
        return raw
    finally:
        workbook.close()


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output', type=Path, default=Path('local-data/game-data.json'))
    parser.add_argument('--report', type=Path, default=Path('local-data/import-report.json'))
    args = parser.parse_args()
    resolved = [path.resolve() for path in (args.workbook, args.output, args.report)]
    if len(set(resolved)) != 3:
        parser.error('Workbook, output, and report must have distinct paths.')
    try:
        data, report = normalize_sheets(read_workbook(args.workbook))
        report['source_sha256'] = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
        write_json(args.output, data)
        write_json(args.report, report)
    except (ValueError, OSError) as exc:
        parser.exit(2, f'Import failed: {exc}\n')
    print(json.dumps({'rows': report['row_counts'], 'price_corrections': len(report['price_corrections']),
        'unknown_prices': len(report['unknown_prices']), 'disabled_recipes': len(report['disabled_recipes'])}, ensure_ascii=False))


if __name__ == '__main__':
    main()
