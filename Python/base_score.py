#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate "料理原始得分表.xlsx"

Behavior (per spec):
- Working directory is set to the folder containing this .py file.
- Reads ../东方夜雀食堂.xlsx
  - Sheet "稀客": columns "Name", "Preferred Tags", "Disliked Tags"
  - Sheet "料理": columns "Name", "Tags"
- For each guest and each recipe, compute:
    score = (# of recipe tags that hit Preferred Tags) - (# of recipe tags that hit Disliked Tags)
  (counted per tag occurrence after splitting; tags are treated as a set per cell after splitting)
- Output columns:
    name, Preferred Tags, Disliked Tags, recipe, Tags, score
- Writes ../料理原始得分表.xlsx
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# Split on common separators. (Keeps tag text itself intact.)
_SPLIT_RE = re.compile(r"[,，;；、\n\r\t|/ ]+")

def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()

def split_tags(v) -> List[str]:
    txt = _s(v)
    if not txt:
        return []
    parts = [p.strip() for p in _SPLIT_RE.split(txt) if p and p.strip()]
    # de-dup while keeping order
    seen = set()
    out = []
    for p in parts:
        if p not in seen:
            out.append(p)
            seen.add(p)
    return out

def sheet_rows_as_dicts(ws) -> List[Dict[str, object]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [(_s(h) if h is not None else "") for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all((c is None or _s(c) == "") for c in r):
            continue
        out.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
    return out

def autosize_columns(ws, max_width: int = 60, min_width: int = 10):
    """Simple autosize based on string length; wraps text for long cells."""
    col_widths = {}
    for col in range(1, ws.max_column + 1):
        best = 0
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            s = str(v)
            # use max line length (newline-aware)
            best = max(best, max((len(line) for line in s.splitlines()), default=0))
        width = min(max(min_width, best + 2), max_width)
        col_widths[col] = width
        ws.column_dimensions[get_column_letter(col)].width = width

    # Center + wrap
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).alignment = align


def main():
    script_dir = Path(__file__).resolve().parent
    os.chdir(script_dir)

    in_path = script_dir.parent / "东方夜雀食堂.xlsx"
    out_path = script_dir.parent / "料理原始得分表.xlsx"

    if not in_path.exists():
        raise FileNotFoundError(f"Input file not found: {in_path}")

    print(f"Working dir: {script_dir}")
    print(f"Input : {in_path}")
    print(f"Output: {out_path}")

    wb = openpyxl.load_workbook(in_path, data_only=True)

    if "稀客" not in wb.sheetnames:
        raise KeyError("Missing sheet: 稀客")
    if "料理" not in wb.sheetnames:
        raise KeyError("Missing sheet: 料理")

    guest_rows = sheet_rows_as_dicts(wb["稀客"])
    dish_rows = sheet_rows_as_dicts(wb["料理"])

    # Parse guests
    guests: List[Tuple[str, str, str, List[str], List[str]]] = []
    # (name, pref_raw, diss_raw, pref_list, diss_list)
    for r in guest_rows:
        name = _s(r.get("Name"))
        if not name:
            continue
        pref_raw = _s(r.get("Preferred Tags"))
        diss_raw = _s(r.get("Disliked Tags"))
        pref_list = split_tags(pref_raw)
        diss_list = split_tags(diss_raw)
        guests.append((name, pref_raw, diss_raw, pref_list, diss_list))

    # Parse dishes
    dishes: List[Tuple[str, str, List[str]]] = []
    # (dish_name, tags_raw, tags_list)
    for r in dish_rows:
        dish_name = _s(r.get("Name"))
        if not dish_name:
            continue
        tags_raw = _s(r.get("Tags"))
        tags_list = split_tags(tags_raw)
        dishes.append((dish_name, tags_raw, tags_list))

    print(f"Parsed: Guests={len(guests)}, Dishes={len(dishes)}")

    # Build output workbook
    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = "Scores"

    headers = ["name", "Preferred Tags", "Disliked Tags", "recipe", "Tags", "score"]
    bold = Font(bold=True)
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = bold

    row_i = 2
    for gi, (gname, pref_raw, diss_raw, pref_list, diss_list) in enumerate(guests, 1):
        pref_set = set(pref_list)
        diss_set = set(diss_list)

        for dish_name, tags_raw, tags_list in dishes:
            score = 0
            for t in tags_list:
                if t in pref_set:
                    score += 1
                if t in diss_set:
                    score -= 1

            ws.cell(row=row_i, column=1, value=gname)
            ws.cell(row=row_i, column=2, value=pref_raw)
            ws.cell(row=row_i, column=3, value=diss_raw)
            ws.cell(row=row_i, column=4, value=dish_name)
            ws.cell(row=row_i, column=5, value=tags_raw)
            ws.cell(row=row_i, column=6, value=score)
            row_i += 1

        if gi % 10 == 0:
            print(f"... processed guests {gi}/{len(guests)}")

    ws.freeze_panes = "A2"
    autosize_columns(ws)

    out_wb.save(out_path)
    print(f"Done. Wrote: {out_path}")


if __name__ == "__main__":
    main()
