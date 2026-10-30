import os
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

from Classes import Food_Class  # 需要你的 Food_Class.py（Food/Drink/Ingredient）


def format_sheet(ws, default_width=31.42, widths=None):
    """设置列宽、字体、居中、全边框。作用范围：当前sheet的已用区域（max_row/max_column）。"""
    widths = widths or {}

    # 1) 列宽（先默认，再按指定列覆盖）
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, default_width)

    # 2) 字体 + 居中
    font = Font(name="等线", size=11)
    align = Alignment(horizontal="center", vertical="center")

    # 3) 全边框（thin）
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = font
            cell.alignment = align
            cell.border = border


def color_column_by_profit_groups(ws, profit_values, target_col, header_row=1):
    """
    按 profit_values（已按写入顺序对应行）分 5 组，
    给 target_col 这一列的单元格填充分组色：高利润更绿，低利润更红。
    """
    n = len(profit_values)
    if n == 0:
        return

    palette_5 = ["FF63BE7B", "FFA9D18E", "FFFFD966", "FFF4B183", "FFE06666"]

    k = min(5, n)
    if k < 5:
        idxs = [round(i * (len(palette_5) - 1) / (k - 1)) if k > 1 else 0 for i in range(k)]
        palette = [palette_5[i] for i in idxs]
    else:
        palette = palette_5

    s = pd.Series(profit_values, dtype="float64")
    rank = s.rank(method="first", ascending=False)  # 1 = 最大利润
    groups = pd.qcut(rank, q=k, labels=False)       # 0 = 最高组, k-1 = 最低组

    for i, g in enumerate(groups):
        row = header_row + 1 + i  # 从第2行开始(跳过表头)
        fill = PatternFill(fill_type="solid", fgColor=palette[int(g)])
        ws.cell(row=row, column=target_col).fill = fill


def highlight_available_name_cells(ws, available_names: set, name_col=1, header_row=1, color="FFB6D7A8"):
    """若名称在 available_names 中，则把该行“名称单元格”填充指定颜色（默认浅绿）。"""
    if not available_names:
        return
    fill = PatternFill(fill_type="solid", fgColor=color)
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=name_col).value
        if v is None:
            continue
        if str(v) in available_names:
            ws.cell(row=r, column=name_col).fill = fill


def get_available_food_names(df_recipe: pd.DataFrame) -> set:
    """
    从“料理”sheet中提取 Owned==1 的 name 集合。（旧函数名保留不改，避免其它地方引用炸）
    - 支持表头“模糊匹配”：Name/NAME/name/名称；Owned/OWNED/owned/持有/available
    """
    if df_recipe is None or df_recipe.empty:
        return set()

    name_col = Food_Class.find_df_col(df_recipe, ["Name", "name", "名称"])
    owned_col = Food_Class.find_df_col(df_recipe, ["Owned", "owned", "持有", "available"])

    if name_col is None or owned_col is None:
        return set()

    s = df_recipe[owned_col]
    # 数字/字符串/空值都兼容：非 0 视作已持有
    owned_num = pd.to_numeric(s, errors="coerce").fillna(0)
    mask = owned_num != 0

    names = df_recipe.loc[mask, name_col].dropna().astype(str)
    return set(names.tolist())

def add_price_tier_and_ranks(df, price_col="售价", profit_col="实际利润", q=5):
    """给df增加：价格档、利润率、档内排名（利润/利润率）。"""
    df2 = df.copy()

    uniq = df2[price_col].nunique(dropna=True)
    bins = max(1, min(q, uniq))
    labels = [f"第{i+1}档" for i in range(bins)]

    df2["价格档"] = pd.qcut(df2[price_col], q=bins, labels=labels, duplicates="drop")
    df2["利润率"] = df2[profit_col] / df2[price_col]

    df2["档内利润排名"] = (
        df2.groupby("价格档")[profit_col]
           .rank(method="first", ascending=False)
           .astype(int)
    )
    df2["档内利润率排名"] = (
        df2.groupby("价格档")["利润率"]
           .rank(method="first", ascending=False)
           .astype(int)
    )
    return df2


def build_full_tier_sheet(df_with_tiers):
    """
    输出全量：按价格档 + 档内利润排名排序。
    自动兼容是否存在“材料”列。

    """
    out = df_with_tiers.sort_values(
        ["价格档", "档内利润排名", "实际利润"],
        ascending=[True, True, False]
    ).reset_index(drop=True)

    cols = ["名称"]
    if "材料" in out.columns:
        cols.append("材料")
    cols += ["售价", "实际利润", "利润率", "价格档", "档内利润排名", "档内利润率排名"]
    return out[cols]


def apply_colorscale(ws, col_letter,
                     start_color="FFE06666", mid_color="FFFFD966", end_color="FF63BE7B"):
    """对某列(从第2行到最后)加三色渐变：红-黄-绿。"""
    if ws.max_row < 2:
        return
    rule = ColorScaleRule(
        start_type="min", start_color=start_color,
        mid_type="percentile", mid_value=50, mid_color=mid_color,
        end_type="max", end_color=end_color
    )
    rng = f"{col_letter}2:{col_letter}{ws.max_row}"
    ws.conditional_formatting.add(rng, rule)


def materials_text(food_obj) -> str:
    """把料理的材料列表拼成字符串；遇到空/任意则按原样处理。"""
    try:
        if not food_obj.required_ingredients:
            return ""
        return "、".join([str(x) for x in food_obj.required_ingredients])
    except Exception:
        return ""


def _excel_text_width(s: str) -> int:
    """估算 Excel 显示宽度：英文/数字算1，中文等全角字符算2。"""
    w = 0
    for ch in s:
        if ch in "\n\r\t":
            continue
        w += 2 if unicodedata.east_asian_width(ch) in ("F", "W") else 1
    return w


def autosize_column(ws, col_idx: int, min_width=10, max_width=80, padding=2):
    """按列内容估算列宽并设置（近似自适应）。"""
    max_w = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col_idx).value
        if v is None:
            continue
        s = str(v)
        max_line = max(s.splitlines(), key=len, default="")
        max_w = max(max_w, _excel_text_width(max_line))

    width = max(min_width, min(max_width, max_w + padding))
    ws.column_dimensions[get_column_letter(col_idx)].width = width


def main():
    BASE_DIR = Path(__file__).resolve().parent
    os.chdir(BASE_DIR)

    input_xlsx = BASE_DIR.parent / "东方夜雀食堂.xlsx"
    output_xlsx = BASE_DIR.parent / "ALL利润.xlsx"

    Food_Dict = {}
    Drink_Dict = {}
    Ingredient_Dict = {}

    sheet_list = pd.read_excel(input_xlsx, sheet_name=None)

    # available==1 的料理名集合（用于名称列高亮）
    available_food_names = get_available_food_names(sheet_list.get("料理"))

    # 实例化对象
    if "食材" in sheet_list:
        for row in sheet_list["食材"].index:
            ing = Food_Class.Ingredient(sheet_list["食材"].loc[row])
            Ingredient_Dict[ing.name] = ing

    if "料理" in sheet_list:
        for row in sheet_list["料理"].index:
            food = Food_Class.Food(sheet_list["料理"].loc[row])
            Food_Dict[food.name] = food

    if "酒水" in sheet_list:
        for row in sheet_list["酒水"].index:
            drink = Food_Class.Drink(sheet_list["酒水"].loc[row])
            Drink_Dict[drink.name] = drink

    # 计算料理实际成本/实际利润（依赖食材dict）
    for food in Food_Dict.values():
        food.after_init(Ingredient_Dict)

    # === 全料理价目 ===
    df_food = pd.DataFrame({
        "名称": [f.name for f in Food_Dict.values()],
        "材料": [materials_text(f) for f in Food_Dict.values()],
        "售价": [f.sell_price for f in Food_Dict.values()],
        "实际利润": [f.profit_actual for f in Food_Dict.values()],
    }).sort_values(by="实际利润", ascending=False)

    # === 全酒水价目 ===
    df_drink = pd.DataFrame({
        "名称": [d.name for d in Drink_Dict.values()],
        "售价": [d.sell_price for d in Drink_Dict.values()],
        "实际利润": [d.profit for d in Drink_Dict.values()],
    }).sort_values(by="实际利润", ascending=False)

    # 列宽：移除序号后，列号整体左移
    # 全料理价目：名称A 材料B 售价C 实际利润D
    col_widths_food_4 = {1: 20.27, 3: 12.18, 4: 12.18}  # 第2列材料不固定
    # 全酒水价目：名称A 售价B 实际利润C
    col_widths_drink_3 = {1: 20.27, 2: 12.18, 3: 12.18}

    # 价格梯度_全量：移除序号后
    # 全料理梯度：名称A 材料B 售价C 实际利润D 利润率E 价格档F 档内利润排名G 档内利润率排名H
    col_widths_food_tier_8 = {1: 20.27, 3: 12.18, 4: 12.18, 5: 12.18, 6: 12.18, 7: 12.18, 8: 14.0}
    # 全酒水梯度：名称A 售价B 实际利润C 利润率D 价格档E 档内利润排名F 档内利润率排名G
    col_widths_drink_tier_7 = {1: 20.27, 2: 12.18, 3: 12.18, 4: 12.18, 5: 12.18, 6: 12.18, 7: 14.0}

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        # === 原两张表 ===
        df_food.to_excel(writer, sheet_name="全料理价目", index=False)
        df_drink.to_excel(writer, sheet_name="全酒水价目", index=False)

        ws_food = writer.book["全料理价目"]
        ws_drink = writer.book["全酒水价目"]

        format_sheet(ws_food, default_width=31.42, widths=col_widths_food_4)
        format_sheet(ws_drink, default_width=31.42, widths=col_widths_drink_3)

        # 材料列（第2列）自适应
        autosize_column(ws_food, 2, min_width=12, max_width=90, padding=2)

        # 全料理价目：名称A 材料B 售价C 实际利润D -> col=4
        color_column_by_profit_groups(ws_food, df_food["实际利润"].tolist(), target_col=4, header_row=1)
        # 全酒水价目：名称A 售价B 实际利润C -> col=3
        color_column_by_profit_groups(ws_drink, df_drink["实际利润"].tolist(), target_col=3, header_row=1)

        # ✅ 名称列只高亮 available==1（全料理价目：名称在第1列）
        highlight_available_name_cells(ws_food, available_food_names, name_col=1, header_row=1, color="FFB6D7A8")

        # === 价格梯度（全量）两张表 ===
        df_food_tier = add_price_tier_and_ranks(df_food, price_col="售价", profit_col="实际利润", q=5)
        df_drink_tier = add_price_tier_and_ranks(df_drink, price_col="售价", profit_col="实际利润", q=5)

        df_food_tier_full = build_full_tier_sheet(df_food_tier)    # 会包含“材料”
        df_drink_tier_full = build_full_tier_sheet(df_drink_tier)  # 不含“材料”

        df_food_tier_full.to_excel(writer, sheet_name="全料理梯度", index=False)
        df_drink_tier_full.to_excel(writer, sheet_name="全酒水梯度", index=False)

        ws_food_tier = writer.book["全料理梯度"]
        ws_drink_tier = writer.book["全酒水梯度"]

        format_sheet(ws_food_tier, default_width=12.18, widths=col_widths_food_tier_8)
        format_sheet(ws_drink_tier, default_width=12.18, widths=col_widths_drink_tier_7)

        # 材料列（第2列）自适应
        autosize_column(ws_food_tier, 2, min_width=12, max_width=90, padding=2)

        # 冻结表头 + 自动筛选
        ws_food_tier.freeze_panes = "A2"
        ws_drink_tier.freeze_panes = "A2"
        ws_food_tier.auto_filter.ref = ws_food_tier.dimensions
        ws_drink_tier.auto_filter.ref = ws_drink_tier.dimensions

        # 渐变色：售价、实际利润、利润率（价格梯度表）
        # 全料理梯度：名称A 材料B 售价C 实际利润D 利润率E ...
        apply_colorscale(ws_food_tier, "C")  # 售价
        apply_colorscale(ws_food_tier, "D")  # 实际利润
        apply_colorscale(ws_food_tier, "E")  # 利润率

        # 全酒水梯度：名称A 售价B 实际利润C 利润率D ...
        apply_colorscale(ws_drink_tier, "B")  # 售价
        apply_colorscale(ws_drink_tier, "C")  # 实际利润
        apply_colorscale(ws_drink_tier, "D")  # 利润率

        highlight_available_name_cells(ws_food_tier, available_food_names, name_col=1, header_row=1, color="FFB6D7A8")

    print(f"已导出：{output_xlsx}")


if __name__ == "__main__":
    main()
